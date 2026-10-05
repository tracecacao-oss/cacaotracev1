"""Lectura de hojas de cálculo: CSV y Excel (.xlsx). Devuelve filas de textos, sin interpretarlas.

La usan las listas de coordenadas de una parcela y la carga masiva de productores.
"""

import csv
import io
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

# Un .xlsx es un zip: se rechaza antes de abrirlo si descomprimido pasa de este tamaño.
MAXIMO_DESCOMPRIMIDO = 50 * 1024 * 1024


class ErrorTabla(Exception):
    def __init__(self, mensaje: str):
        super().__init__(mensaje)
        self.mensaje = mensaje


def comprobar_zip(contenido: bytes, maximo: int = MAXIMO_DESCOMPRIMIDO) -> zipfile.ZipFile:
    """Abre un zip y se niega si declara más de `maximo` bytes descomprimidos o demasiadas entradas."""
    try:
        archivo = zipfile.ZipFile(io.BytesIO(contenido))
    except zipfile.BadZipFile as exc:
        raise ErrorTabla("El archivo comprimido está dañado o no es un .zip válido.") from exc
    entradas = archivo.infolist()
    if len(entradas) > 500 or sum(e.file_size for e in entradas) > maximo:
        raise ErrorTabla("El archivo comprimido es demasiado grande al descomprimirlo.")
    return archivo


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "sí" if valor else "no"
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, datetime | date):
        return valor.isoformat()
    return str(valor).strip()


def leer_csv(contenido: bytes) -> list[list[str]]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel en Windows guarda los CSV en esta codificación.
        texto = contenido.decode("cp1252", errors="replace")
    muestra = texto[:4096]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    return [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), dialecto)]


def leer_xlsx(contenido: bytes) -> list[list[str]]:
    comprobar_zip(contenido)
    try:
        # Con defusedxml instalado, openpyxl lo usa para leer el XML del libro.
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl lanza varios tipos según el daño del archivo
        raise ErrorTabla("No se pudo leer el Excel. Guárdalo de nuevo como .xlsx o como .csv.") from exc
    try:
        hoja = libro.worksheets[0]
        return [[_texto(v) for v in fila] for fila in hoja.iter_rows(values_only=True)]
    finally:
        libro.close()


def leer_tabla(extension: str, contenido: bytes) -> list[list[str]]:
    """Filas de un .csv o un .xlsx sin las celdas vacías del final. Una fila vacía queda como []."""
    filas = leer_xlsx(contenido) if extension == ".xlsx" else leer_csv(contenido)
    limpias = []
    for fila in filas:
        while fila and not fila[-1]:
            fila = fila[:-1]
        limpias.append(fila)
    while limpias and not limpias[-1]:
        limpias.pop()
    return limpias
